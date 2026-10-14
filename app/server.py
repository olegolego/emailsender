"""Local web app + JSON API. Listens on 127.0.0.1 only."""
import csv
import io
import json
import os
import random
import re
from datetime import timedelta

from flask import Flask, Response, abort, jsonify, request, send_file, send_from_directory

from . import analyzer, db, engine, keepa, mailer, schedule, secrets, sequences, stats, templating, worker

STATIC = db.BASE_DIR / "static"
app = Flask(__name__, static_folder=None)


def body():
    return request.get_json(silent=True) or {}


def ok(**kw):
    return jsonify({"ok": True, **kw})


@app.before_request
def guard():
    # Only this machine, only our own page: blocks other websites and DNS-rebinding tricks.
    host = (request.host or "").rsplit(":", 1)[0].lower()
    if host not in ("127.0.0.1", "localhost"):
        abort(403)
    if request.path.startswith("/api/") and request.method != "GET" and request.headers.get("X-EmailManager") != "1":
        abort(403)


@app.teardown_request
def close_db(exc):
    db.close()


@app.errorhandler(ValueError)
@app.errorhandler(RuntimeError)
@app.errorhandler(mailer.MailError)
@app.errorhandler(keepa.KeepaError)
def value_error(e):
    return jsonify({"ok": False, "error": str(e)}), 400


@app.errorhandler(Exception)
def any_error(e):
    from werkzeug.exceptions import HTTPException
    if isinstance(e, HTTPException):
        return e
    import traceback
    traceback.print_exc()
    return jsonify({"ok": False, "error": f"{type(e).__name__}: {e}"}), 500


# ---------------------------------------------------------------- static

@app.get("/")
def index():
    return send_from_directory(STATIC, "index.html")


@app.get("/static/<path:p>")
def static_files(p):
    return send_from_directory(STATIC, p)


# ---------------------------------------------------------------- dashboard / state

@app.get("/api/state")
def state():
    start, end = schedule.local_day_bounds()
    accounts = db.q("SELECT id, email, daily_limit, active FROM accounts")
    for a in accounts:
        a["sent_today"] = engine.sent_today(a["id"])
    nxt = db.q("SELECT q.id, q.to_addr, q.subject, q.scheduled_at, q.status, q.node_id, p.name AS pipeline, "
               "c.first_name, c.last_name, co.name AS company FROM queue q LEFT JOIN pipelines p ON p.id=q.pipeline_id "
               "LEFT JOIN contacts c ON c.id=q.contact_id LEFT JOIN companies co ON co.id=c.company_id "
               "WHERE q.status IN ('pending','needs_approval') ORDER BY q.status='pending' DESC, q.scheduled_at LIMIT 8")
    recent = db.q("SELECT m.id, m.from_addr, m.from_name, m.subject, m.summary, m.category, m.interest_score, m.ts, m.handled, "
                  "co.name AS company FROM messages m LEFT JOIN companies co ON co.id=m.company_id "
                  "WHERE m.direction='in' ORDER BY m.ts DESC LIMIT 8")
    s7 = stats.summary(days=7)
    return jsonify({
        "worker": worker.status,
        "settings": {k: db.get_setting(k) for k in ("test_mode", "test_address", "sending_paused", "timezone")},
        "ai": analyzer.ai_available(),
        "accounts": accounts,
        "counts": {
            "pending": db.scalar("SELECT COUNT(*) FROM queue WHERE status='pending'"),
            "needs_approval": db.scalar("SELECT COUNT(*) FROM queue WHERE status='needs_approval'"),
            "failed": db.scalar("SELECT COUNT(*) FROM queue WHERE status='failed'"),
            "due_24h": db.scalar("SELECT COUNT(*) FROM queue WHERE status='pending' AND scheduled_at<=?",
                                 (db.iso(db.utcnow() + timedelta(hours=24)),)),
            "unhandled": db.scalar("SELECT COUNT(*) FROM messages WHERE direction='in' AND handled=0"),
            "active_enrollments": db.scalar("SELECT COUNT(*) FROM enrollments WHERE state='active'"),
            "contacts": db.scalar("SELECT COUNT(*) FROM contacts"),
            "sent_today": db.scalar("SELECT COUNT(*) FROM messages WHERE direction='out' AND ts>=? AND ts<?",
                                    (db.iso(start), db.iso(end))),
            "unread_notifications": db.scalar("SELECT COUNT(*) FROM notifications WHERE read=0"),
            "new_price_lists": db.scalar("SELECT COUNT(*) FROM price_lists WHERE status='new'"),
        },
        "week": s7, "next": nxt, "recent": recent, "now": db.now_iso(),
    })


@app.post("/api/worker/check-inbox")
def check_inbox_now():
    worker.request_inbox_check()
    return ok()


@app.post("/api/worker/pause")
def pause_sending():
    db.set_setting("sending_paused", bool(body().get("paused")))
    worker.wake()
    return ok(paused=db.get_setting("sending_paused"))


# ---------------------------------------------------------------- settings & accounts

SETTING_KEYS = set(db.DEFAULT_SETTINGS)


@app.get("/api/settings")
def get_settings():
    s = db.all_settings()
    s["anthropic_key_set"] = secrets.has_secret("anthropic_api_key")
    s["keepa_key_set"] = secrets.has_secret("keepa_api_key")
    s["ai_available"] = analyzer.ai_available()
    return jsonify(s)


@app.post("/api/settings")
def save_settings():
    d = body()
    for k, v in d.items():
        if k in SETTING_KEYS:
            db.set_setting(k, v)
    if "anthropic_api_key" in d:
        secrets.set_secret("anthropic_api_key", (d["anthropic_api_key"] or "").strip())
    if "keepa_api_key" in d:
        secrets.set_secret("keepa_api_key", (d["keepa_api_key"] or "").strip())
    worker.wake()
    return ok()


@app.post("/api/settings/test-ai")
def test_ai():
    if not analyzer.ai_available():
        return jsonify({"ok": False, "error": "No Anthropic API key saved."}), 400
    sample = {"id": 0, "contact_id": None, "from_addr": "jane@example-distributor.com", "from_name": "Jane",
              "subject": "Re: Wholesale account inquiry", "ts": db.now_iso(), "is_bounce": 0,
              "reply_text": "Hi Oleg, thanks for reaching out. Please fill out our credit application at the link "
                            "below and send your resale certificate. What volume do you expect per month?"}
    a, src = analyzer.analyze(sample, use_ai=True)
    return ok(analysis=a, source=src)


ACCOUNT_FIELDS = ("label", "email", "from_name", "username", "smtp_host", "smtp_port", "smtp_security", "imap_host",
                  "imap_port", "imap_folders", "daily_limit", "min_gap_minutes", "signature", "active")


@app.get("/api/accounts")
def list_accounts():
    rows = db.q("SELECT * FROM accounts ORDER BY id")
    for r in rows:
        r["password_set"] = secrets.has_secret(f"account:{r['id']}")
        r["sent_today"] = engine.sent_today(r["id"])
    return jsonify(rows)


@app.post("/api/accounts")
def create_account():
    d = body()
    if "@" not in (d.get("email") or ""):
        raise ValueError("Enter the email address to send from.")
    row = {k: d[k] for k in ACCOUNT_FIELDS if k in d}
    row["email"] = row["email"].strip().lower()
    row["created_at"] = db.now_iso()
    aid = db.insert("accounts", row)
    if d.get("password"):
        secrets.set_secret(f"account:{aid}", d["password"].replace(" ", ""))
    return ok(id=aid)


@app.put("/api/accounts/<int:aid>")
def update_account(aid):
    d = body()
    db.update("accounts", aid, {k: d[k] for k in ACCOUNT_FIELDS if k in d})
    engine.space_out(aid)
    if d.get("password"):
        secrets.set_secret(f"account:{aid}", d["password"].replace(" ", ""))
    return ok()


@app.delete("/api/accounts/<int:aid>")
def delete_account(aid):
    if db.scalar("SELECT COUNT(*) FROM messages WHERE account_id=?", (aid,)):
        db.update("accounts", aid, {"active": 0})
        return ok(note="Mailbox has history, so it was deactivated instead of deleted.")
    db.execute("DELETE FROM accounts WHERE id=?", (aid,))
    secrets.set_secret(f"account:{aid}", "")
    return ok()


@app.post("/api/accounts/<int:aid>/test")
def test_account(aid):
    acc = db.q1("SELECT * FROM accounts WHERE id=?", (aid,))
    result = {"smtp": None, "imap": None}
    try:
        mailer.test_smtp(acc)
        result["smtp"] = "ok"
    except Exception as e:
        result["smtp"] = str(e)
    try:
        n = mailer.test_imap(acc)
        result["imap"] = f"ok ({n} messages in INBOX)"
    except Exception as e:
        result["imap"] = str(e)
    return jsonify(result)


# ---------------------------------------------------------------- contacts & companies

def _contact_rows(where="1=1", args=(), limit=500, offset=0):
    return db.q(f"SELECT c.*, co.name AS company, co.domain AS company_domain, co.category AS company_category, "
                f"(SELECT p.name FROM enrollments e JOIN pipelines p ON p.id=e.pipeline_id WHERE e.contact_id=c.id "
                f" ORDER BY e.id DESC LIMIT 1) AS pipeline, "
                f"(SELECT e.state FROM enrollments e WHERE e.contact_id=c.id ORDER BY e.id DESC LIMIT 1) AS pipeline_state, "
                f"(SELECT COUNT(*) FROM messages m WHERE m.contact_id=c.id AND m.direction='out') AS sent_count, "
                f"(SELECT COUNT(*) FROM messages m WHERE m.contact_id=c.id AND m.direction='in') AS reply_count "
                f"FROM contacts c LEFT JOIN companies co ON co.id=c.company_id WHERE {where} "
                f"ORDER BY c.id DESC LIMIT ? OFFSET ?", args + (limit, offset))


@app.get("/api/contacts")
def list_contacts():
    where, args = ["1=1"], []
    s = (request.args.get("q") or "").strip()
    if s:
        like = f"%{s.lower()}%"
        where.append("(lower(c.email) LIKE ? OR lower(c.first_name||' '||c.last_name) LIKE ? OR lower(co.name) LIKE ? "
                     "OR lower(c.tags) LIKE ?)")
        args += [like] * 4
    if request.args.get("status"):
        where.append("c.status=?")
        args.append(request.args["status"])
    if request.args.get("tag"):
        where.append("(','||lower(c.tags)||',') LIKE ?")
        args.append(f"%,{request.args['tag'].lower()},%")
    if request.args.get("company_id"):
        where.append("c.company_id=?")
        args.append(int(request.args["company_id"]))
    if request.args.get("not_in_pipeline"):
        where.append("NOT EXISTS (SELECT 1 FROM enrollments e WHERE e.contact_id=c.id AND e.pipeline_id=?)")
        args.append(int(request.args["not_in_pipeline"]))
    w = " AND ".join(where)
    total = db.scalar(f"SELECT COUNT(*) FROM contacts c LEFT JOIN companies co ON co.id=c.company_id WHERE {w}", tuple(args))
    rows = _contact_rows(w, tuple(args), int(request.args.get("limit", 500)), int(request.args.get("offset", 0)))
    tags = sorted({t.strip() for r in db.q("SELECT tags FROM contacts WHERE tags<>''") for t in r["tags"].split(",") if t.strip()})
    return jsonify({"total": total, "rows": rows, "tags": tags})


@app.post("/api/contacts")
def create_contact():
    d = body()
    cid = engine.upsert_contact(d)
    return ok(id=cid)


@app.get("/api/contacts/<int:cid>")
def get_contact(cid):
    rows = _contact_rows("c.id=?", (cid,))
    if not rows:
        abort(404)
    c = rows[0]
    c["custom"] = json.loads(c.get("custom_json") or "{}")
    c["messages"] = db.q("SELECT id, direction, subject, body, reply_text, ts, category, sentiment, interest_score, summary, "
                         "is_test, is_bounce, from_addr, to_addr, node_id FROM messages WHERE contact_id=? ORDER BY ts", (cid,))
    c["enrollments"] = db.q("SELECT e.*, p.name AS pipeline FROM enrollments e JOIN pipelines p ON p.id=e.pipeline_id "
                            "WHERE e.contact_id=? ORDER BY e.id DESC", (cid,))
    for e in c["enrollments"]:
        e["path"] = json.loads(e["path"] or "[]")
    c["queue"] = db.q("SELECT * FROM queue WHERE contact_id=? AND status IN ('pending','needs_approval') ORDER BY scheduled_at", (cid,))
    c["events"] = db.q("SELECT * FROM events WHERE contact_id=? ORDER BY ts DESC LIMIT 100", (cid,))
    return jsonify(c)


@app.put("/api/contacts/<int:cid>")
def update_contact(cid):
    d = body()
    fields = {k: d[k] for k in ("first_name", "last_name", "title", "phone", "timezone", "tags", "status", "notes") if k in d}
    if "custom" in d:
        fields["custom_json"] = json.dumps(d["custom"])
    if "company" in d:
        fields["company_id"] = engine.upsert_company(d.get("company") or "", d.get("company_domain") or "",
                                                     d.get("website") or "", d.get("category") or "") if d["company"] else None
    if d.get("email"):
        fields["email"] = d["email"].strip().lower()
    db.update("contacts", cid, fields)
    if fields.get("status") in engine.SUPPRESSED:
        for e in db.q("SELECT * FROM enrollments WHERE contact_id=? AND state IN ('active','paused')", (cid,)):
            engine.finish(e, "stop", f"contact marked {fields['status']}")
    return ok()


@app.delete("/api/contacts/<int:cid>")
def delete_contact(cid):
    for e in db.q("SELECT * FROM enrollments WHERE contact_id=? AND state IN ('active','paused')", (cid,)):
        engine.finish(e, "stop", "contact deleted")
    with db.transaction() as cx:
        cx.execute("DELETE FROM queue WHERE contact_id=?", (cid,))
        cx.execute("UPDATE messages SET contact_id=NULL WHERE contact_id=?", (cid,))
        cx.execute("UPDATE events SET contact_id=NULL WHERE contact_id=?", (cid,))
        cx.execute("DELETE FROM enrollments WHERE contact_id=?", (cid,))
        cx.execute("DELETE FROM contacts WHERE id=?", (cid,))
    return ok()


@app.post("/api/contacts/bulk")
def bulk_contacts():
    d = body()
    ids = [int(i) for i in d.get("ids") or []]
    action = d.get("action")
    for cid in ids:
        if action == "tag":
            c = db.q1("SELECT tags FROM contacts WHERE id=?", (cid,))
            tags = {t.strip() for t in ((c["tags"] if c else "") + "," + d.get("value", "")).split(",") if t.strip()}
            db.update("contacts", cid, {"tags": ",".join(sorted(tags))})
        elif action == "status":
            db.update("contacts", cid, {"status": d.get("value")})
            if d.get("value") in engine.SUPPRESSED:
                for e in db.q("SELECT * FROM enrollments WHERE contact_id=? AND state IN ('active','paused')", (cid,)):
                    engine.finish(e, "stop", f"contact marked {d.get('value')}")
        elif action == "delete":
            delete_contact(cid)
    return ok(count=len(ids))


FIELD_ALIASES = {
    "email": ["email", "e-mail", "email address", "mail", "contact email"],
    "first_name": ["first_name", "first name", "firstname", "first", "given name"],
    "last_name": ["last_name", "last name", "lastname", "last", "surname", "family name"],
    "full_name": ["name", "full name", "contact", "contact name"],
    "company": ["company", "company name", "organization", "organisation", "distributor", "brand", "business"],
    "company_domain": ["domain", "company domain"],
    "website": ["website", "url", "site", "web"],
    "title": ["title", "job title", "position", "role"],
    "phone": ["phone", "phone number", "telephone", "tel"],
    "category": ["category", "industry", "niche", "vertical"],
    "tags": ["tags", "tag", "list", "segment"],
    "timezone": ["timezone", "time zone", "tz"],
    "notes": ["notes", "note", "comments"],
}


@app.post("/api/contacts/import")
def import_contacts():
    d = body()
    text = d.get("csv") or ""
    extra_tags = (d.get("tags") or "").strip()
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except Exception:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    headers = reader.fieldnames or []
    mapping = {}
    for h in headers:
        key = h.strip().lower()
        for field, names in FIELD_ALIASES.items():
            if key in names and field not in mapping.values():
                mapping[h] = field
                break
    if "email" not in mapping.values():
        raise ValueError(f"No email column found. Columns: {', '.join(headers)}")
    created, updated, errors = 0, 0, []
    for i, row in enumerate(reader, start=2):
        rec, custom = {}, {}
        for h, v in row.items():
            if h is None:
                continue
            v = (v or "").strip()
            if h in mapping:
                rec[mapping[h]] = v
            elif v:
                custom[h.strip().lower().replace(" ", "_")] = v
        if rec.get("full_name") and not rec.get("first_name"):
            parts = rec["full_name"].split()
            rec["first_name"], rec["last_name"] = parts[0], " ".join(parts[1:])
        emails = mailer.EMAIL_RE.findall(rec.get("email") or "")
        if not emails:
            if any(row.values()):
                errors.append(f"row {i}: no valid email")
            continue
        rec["email"] = emails[0]
        rec["custom"] = custom
        if extra_tags:
            rec["tags"] = ",".join(filter(None, [rec.get("tags", ""), extra_tags]))
        existed = db.q1("SELECT id FROM contacts WHERE email=?", (rec["email"].lower(),))
        try:
            engine.upsert_contact(rec)
            updated += 1 if existed else 0
            created += 0 if existed else 1
        except Exception as e:
            errors.append(f"row {i}: {e}")
    return ok(created=created, updated=updated, errors=errors[:50], mapping=mapping)


@app.get("/api/companies")
def list_companies():
    return jsonify(db.q("SELECT co.*, (SELECT COUNT(*) FROM contacts c WHERE c.company_id=co.id) AS contacts FROM companies co "
                        "ORDER BY co.name"))


# ---------------------------------------------------------------- pipelines

@app.get("/api/pipelines")
def list_pipelines():
    return jsonify(stats.pipeline_overview())


@app.post("/api/pipelines")
def create_pipeline():
    d = body()
    defn = d.get("definition") or {"start": "s1", "nodes": [
        {"id": "s1", "type": "email", "name": "First email", "thread": "new", "subject": "", "body": "",
         "delay": {"value": 0, "unit": "minutes"}, "wait": {"value": 3, "unit": "days"}, "branches": {}}]}
    pid = db.insert("pipelines", {"name": d.get("name") or "New pipeline", "description": d.get("description") or "",
                                  "definition": json.dumps(defn),
                                  "settings": json.dumps(d.get("settings") or schedule.DEFAULT_PIPELINE_SETTINGS),
                                  "status": "draft", "account_id": d.get("account_id"),
                                  "created_at": db.now_iso(), "updated_at": db.now_iso()})
    return ok(id=pid)


@app.get("/api/pipelines/<int:pid>")
def get_pipeline(pid):
    p = engine.load_pipeline(pid)
    if not p:
        abort(404)
    p.pop("nodes", None)
    p["steps"] = stats.by_step(pid)
    p["errors"] = engine.validate_definition(p["definition"])
    return jsonify(p)


@app.put("/api/pipelines/<int:pid>")
def save_pipeline(pid):
    d = body()
    upd = {"updated_at": db.now_iso()}
    for k in ("name", "description", "account_id"):
        if k in d:
            upd[k] = d[k]
    if "definition" in d:
        upd["definition"] = json.dumps(d["definition"])
    if "settings" in d:
        upd["settings"] = json.dumps(d["settings"])
    db.update("pipelines", pid, upd)
    errs = engine.validate_definition(d["definition"]) if "definition" in d else []
    if "definition" in d:
        ids = {n["id"] for n in d["definition"].get("nodes", [])}
        for e in db.q("SELECT * FROM enrollments WHERE pipeline_id=? AND state IN ('active','paused')", (pid,)):
            if e["current_node"] not in ids:
                engine.finish(e, "stop", "its step was deleted from the pipeline")
    return ok(errors=errs)


@app.post("/api/pipelines/<int:pid>/status")
def pipeline_status(pid):
    st = body().get("status")
    if st not in ("draft", "active", "paused", "archived"):
        raise ValueError("bad status")
    if st == "active":
        p = engine.load_pipeline(pid)
        errs = engine.validate_definition(p["definition"])
        if errs:
            raise ValueError(" ".join(errs))
    db.update("pipelines", pid, {"status": st, "updated_at": db.now_iso()})
    worker.wake()
    return ok()


@app.post("/api/pipelines/<int:pid>/duplicate")
def duplicate_pipeline(pid):
    p = db.q1("SELECT * FROM pipelines WHERE id=?", (pid,))
    nid = db.insert("pipelines", {"name": p["name"] + " (copy)", "description": p["description"],
                                  "definition": p["definition"], "settings": p["settings"], "status": "draft",
                                  "account_id": p["account_id"], "created_at": db.now_iso(), "updated_at": db.now_iso()})
    return ok(id=nid)


@app.delete("/api/pipelines/<int:pid>")
def delete_pipeline(pid):
    if db.scalar("SELECT COUNT(*) FROM messages WHERE pipeline_id=?", (pid,)):
        db.update("pipelines", pid, {"status": "archived"})
        for e in db.q("SELECT * FROM enrollments WHERE pipeline_id=? AND state IN ('active','paused')", (pid,)):
            engine.finish(e, "stop", "pipeline archived")
        return ok(note="Pipeline has sent emails, so it was archived to keep the statistics.")
    with db.transaction() as cx:
        cx.execute("DELETE FROM queue WHERE pipeline_id=?", (pid,))
        cx.execute("UPDATE events SET enrollment_id=NULL WHERE enrollment_id IN (SELECT id FROM enrollments WHERE pipeline_id=?)", (pid,))
        cx.execute("DELETE FROM enrollments WHERE pipeline_id=?", (pid,))
        cx.execute("DELETE FROM pipelines WHERE id=?", (pid,))
    return ok()


@app.post("/api/pipelines/<int:pid>/enroll")
def enroll(pid):
    d = body()
    ids = [int(i) for i in d.get("contact_ids") or []]
    if d.get("tag"):
        ids += [r["id"] for r in db.q("SELECT id FROM contacts WHERE (','||lower(tags)||',') LIKE ?",
                                      (f"%,{d['tag'].lower()},%",))]
    if d.get("all_new"):
        ids += [r["id"] for r in db.q("SELECT id FROM contacts WHERE status='new'")]
    ids = list(dict.fromkeys(ids))
    res = engine.enroll(pid, ids, d.get("start_at"), float(d.get("stagger_minutes") or 0), d.get("account_id"))
    worker.wake()
    return ok(**res)


@app.post("/api/pipelines/<int:pid>/preview")
def preview(pid):
    """Render a step for a real contact without sending anything."""
    d = body()
    p = engine.load_pipeline(pid)
    if d.get("definition"):
        p["definition"] = d["definition"]
        p["nodes"] = {n["id"]: n for n in d["definition"].get("nodes", [])}
    node = p["nodes"].get(d.get("node_id"))
    if not node:
        raise ValueError("Unknown step")
    cid = d.get("contact_id") or db.scalar("SELECT id FROM contacts ORDER BY id DESC LIMIT 1")
    if not cid:
        raise ValueError("Add a contact first to preview.")
    fake = {"id": 0, "contact_id": cid, "account_id": engine.pick_account_id(p), "last_subject": "(previous subject)",
            "last_message_id": "<previous@message>" if node.get("thread") == "reply" else "", "references_hdr": ""}
    if node.get("type") == "ai_email" and not d.get("run_ai"):
        c = db.q1("SELECT * FROM contacts WHERE id=?", (cid,))
        return ok(subject="(written by Claude when the step runs)", body="Instructions:\n" + (node.get("prompt") or ""),
                  missing=[], to=c["email"], error="")
    r = engine.render_step(fake, p, node, random.Random())
    times = []
    base = db.utcnow()
    for _ in range(3):
        times.append(db.iso(schedule.plan_send_time(base, node.get("delay"), p["settings"],
                                                    db.q1("SELECT * FROM contacts WHERE id=?", (cid,)))))
    return ok(**r, example_send_times=times)


@app.get("/api/enrollments")
def list_enrollments():
    pid = request.args.get("pipeline_id")
    where, args = "1=1", ()
    if pid:
        where, args = "e.pipeline_id=?", (int(pid),)
    if request.args.get("state"):
        where += " AND e.state=?"
        args += (request.args["state"],)
    rows = db.q(f"SELECT e.*, c.email, c.first_name, c.last_name, c.status AS contact_status, co.name AS company, "
                f"p.name AS pipeline FROM enrollments e JOIN contacts c ON c.id=e.contact_id "
                f"LEFT JOIN companies co ON co.id=c.company_id JOIN pipelines p ON p.id=e.pipeline_id "
                f"WHERE {where} ORDER BY e.updated_at DESC LIMIT 1000", args)
    for r in rows:
        r["path"] = json.loads(r["path"] or "[]")
    return jsonify(rows)


@app.post("/api/enrollments/<int:eid>/action")
def enrollment_action(eid):
    d = body()
    e = db.q1("SELECT * FROM enrollments WHERE id=?", (eid,))
    if not e:
        abort(404)
    a = d.get("action")
    if a == "pause":
        db.update("enrollments", eid, {"state": "paused", "updated_at": db.now_iso()})
        db.execute("UPDATE queue SET status='needs_approval', approval_reason='enrollment paused' "
                   "WHERE enrollment_id=? AND status='pending'", (eid,))
    elif a == "resume":
        upd = {"state": "active", "stopped_reason": "", "updated_at": db.now_iso()}
        if e["phase"] in ("done",):
            upd.update({"phase": "enter", "next_action_at": db.now_iso(), "node_entered_at": db.now_iso()})
        elif e["phase"] != "queued" and not e["next_action_at"]:
            upd["next_action_at"] = db.now_iso()
        db.update("enrollments", eid, upd)
        db.execute("UPDATE queue SET status='pending', approval_reason='' WHERE enrollment_id=? AND status='needs_approval' "
                   "AND approval_reason IN ('batch paused','enrollment paused')", (eid,))
        if e["phase"] == "queued" and not db.scalar(
                "SELECT COUNT(*) FROM queue WHERE enrollment_id=? AND status IN ('pending','needs_approval','sending')", (eid,)):
            db.update("enrollments", eid, {"phase": "enter", "next_action_at": db.now_iso(), "node_entered_at": db.now_iso()})
    elif a == "stop":
        engine.finish(e, "stop", "stopped by me")
    elif a == "move":
        engine.cancel_pending(eid, "moved to another step")
        db.update("enrollments", eid, {"state": "active"})
        engine.goto(db.q1("SELECT * FROM enrollments WHERE id=?", (eid,)), d.get("node_id"), "moved by me")
    else:
        raise ValueError("unknown action")
    worker.wake()
    return ok()


# ---------------------------------------------------------------- queue

@app.get("/api/queue")
def list_queue():
    st = request.args.get("status") or "pending"
    statuses = {"upcoming": ("pending", "needs_approval", "sending")}.get(st, (st,))
    ph = ",".join("?" for _ in statuses)
    order = "q.sent_at DESC" if st in ("sent",) else "q.created_at DESC" if st in ("failed", "cancelled") else "q.scheduled_at"
    rows = db.q(f"SELECT q.*, p.name AS pipeline, p.status AS pipeline_status, c.first_name, c.last_name, c.email, "
                f"co.name AS company, a.email AS account_email, b.name AS batch, b.id AS batch_id FROM queue q "
                f"LEFT JOIN pipelines p ON p.id=q.pipeline_id LEFT JOIN contacts c ON c.id=q.contact_id "
                f"LEFT JOIN companies co ON co.id=c.company_id LEFT JOIN accounts a ON a.id=q.account_id "
                f"LEFT JOIN enrollments e ON e.id=q.enrollment_id LEFT JOIN batches b ON b.id=e.batch_id "
                f"WHERE q.status IN ({ph}) ORDER BY {order} LIMIT 300", statuses)
    for r in rows:
        if r["pipeline_id"]:
            p = engine.load_pipeline(r["pipeline_id"])
            r["step_name"] = (p["nodes"].get(r["node_id"]) or {}).get("name") or r["node_id"]
    counts = {r["status"]: r["n"] for r in db.q("SELECT status, COUNT(*) AS n FROM queue GROUP BY status")}
    return jsonify({"rows": rows, "counts": counts})


@app.post("/api/queue/<int:qid>/action")
def queue_action(qid):
    d = body()
    a = d.get("action")
    item = db.q1("SELECT * FROM queue WHERE id=?", (qid,))
    if not item:
        abort(404)
    if a == "send_now":
        if item["status"] == "needs_approval":
            db.update("queue", qid, {"status": "pending"})
        ok_ = engine.send_item(qid, force=True)
        if not ok_:
            return jsonify({"ok": False, "error": db.scalar("SELECT error FROM queue WHERE id=?", (qid,)) or "Not sent"}), 400
    elif a == "approve":
        db.update("queue", qid, {"status": "pending", "approval_reason": ""})
        engine.space_out(item["account_id"])
        worker.wake()
    elif a == "cancel":
        db.update("queue", qid, {"status": "cancelled", "error": "skipped by me"})
        if item["enrollment_id"] and d.get("advance", True):
            e = db.q1("SELECT * FROM enrollments WHERE id=?", (item["enrollment_id"],))
            if e and e["state"] == "active" and e["phase"] == "queued" and e["current_node"] == item["node_id"]:
                db.update("enrollments", e["id"], {"phase": "waiting", "next_action_at": db.now_iso()})
                worker.wake()
    elif a == "edit":
        upd = {k: d[k] for k in ("subject", "body", "to_addr") if k in d}
        if d.get("scheduled_at"):
            upd["scheduled_at"] = db.iso(db.parse_iso(d["scheduled_at"]))
        db.update("queue", qid, upd)
    elif a == "reschedule":
        db.update("queue", qid, {"scheduled_at": db.iso(db.parse_iso(d["scheduled_at"])),
                                 "status": "pending" if item["status"] in ("failed",) else item["status"]})
    elif a == "retry":
        db.update("queue", qid, {"status": "pending", "scheduled_at": db.now_iso(), "error": ""})
        worker.wake()
    else:
        raise ValueError("unknown action")
    return ok()


@app.post("/api/queue/bulk")
def queue_bulk():
    d = body()
    for qid in d.get("ids") or []:
        if d.get("action") == "approve":
            db.execute("UPDATE queue SET status='pending', approval_reason='' WHERE id=? AND status='needs_approval'", (qid,))
        elif d.get("action") == "cancel":
            db.execute("UPDATE queue SET status='cancelled', error='skipped by me' WHERE id=? AND status IN ('pending','needs_approval')",
                       (qid,))
        elif d.get("action") == "shift":
            mins = float(d.get("minutes") or 0)
            it = db.q1("SELECT scheduled_at FROM queue WHERE id=?", (qid,))
            if it:
                db.update("queue", qid, {"scheduled_at": db.iso(db.parse_iso(it["scheduled_at"]) + timedelta(minutes=mins))})
    worker.wake()
    return ok()


# ---------------------------------------------------------------- inbox / messages

@app.get("/api/inbox")
def inbox():
    f = request.args.get("filter") or "all"
    where = "m.direction='in'"
    if f == "unhandled":
        where += " AND m.handled=0"
    elif f == "positive":
        where += " AND m.category IN ('interested','question','referral')"
    elif f in ("not_interested", "out_of_office", "bounce", "unsubscribe", "auto_reply", "other", "interested",
               "question", "referral"):
        where += f" AND m.category='{f}'"
    rows = db.q(f"SELECT m.id, m.from_addr, m.from_name, m.subject, m.summary, m.category, m.sentiment, m.interest_score, "
                f"m.ts, m.handled, m.contact_id, m.pipeline_id, m.analysis_source, m.replied_to_id, co.name AS company, p.name AS pipeline, "
                f"substr(m.reply_text,1,240) AS snippet FROM messages m LEFT JOIN companies co ON co.id=m.company_id "
                f"LEFT JOIN pipelines p ON p.id=m.pipeline_id WHERE {where} ORDER BY m.ts DESC LIMIT 500")
    counts = {r["k"]: r["n"] for r in db.q("SELECT CASE WHEN handled=0 THEN 'unhandled' ELSE 'handled' END AS k, COUNT(*) AS n "
                                           "FROM messages WHERE direction='in' GROUP BY 1")}
    return jsonify({"rows": rows, "counts": counts})


@app.get("/api/messages/<int:mid>")
def get_message(mid):
    m = db.q1("SELECT * FROM messages WHERE id=?", (mid,))
    if not m:
        abort(404)
    m["analysis"] = json.loads(m["analysis_json"]) if m["analysis_json"] else None
    m["thread"] = db.q("SELECT id, direction, subject, body, reply_text, ts, category, from_addr, to_addr, is_test "
                       "FROM messages WHERE contact_id=? ORDER BY ts", (m["contact_id"],)) if m["contact_id"] else [m]
    for t in m["thread"]:
        t["files"] = db.q("SELECT id, kind, filename, url, size FROM price_lists WHERE message_row_id=?", (t["id"],))
    m["contact"] = db.q1("SELECT c.*, co.name AS company FROM contacts c LEFT JOIN companies co ON co.id=c.company_id "
                         "WHERE c.id=?", (m["contact_id"],)) if m["contact_id"] else None
    m["enrollment"] = db.q1("SELECT e.*, p.name AS pipeline FROM enrollments e JOIN pipelines p ON p.id=e.pipeline_id "
                            "WHERE e.id=?", (m["enrollment_id"],)) if m["enrollment_id"] else None
    if m["enrollment"]:
        m["enrollment"]["path"] = json.loads(m["enrollment"]["path"] or "[]")
        p = engine.load_pipeline(m["enrollment"]["pipeline_id"])
        m["enrollment"]["nodes"] = [{"id": n["id"], "name": n.get("name") or n["id"]} for n in p["definition"]["nodes"]]
    return jsonify(m)


@app.post("/api/messages/<int:mid>/handled")
def mark_handled(mid):
    db.update("messages", mid, {"handled": 1 if body().get("handled", True) else 0})
    return ok()


@app.post("/api/messages/<int:mid>/reclassify")
def reclassify(mid):
    d = body()
    cat = d.get("category")
    if cat not in analyzer.CATEGORIES:
        raise ValueError("bad category")
    m = db.q1("SELECT * FROM messages WHERE id=?", (mid,))
    a = json.loads(m["analysis_json"]) if m["analysis_json"] else analyzer.rules(m)
    a["category"] = cat
    db.update("messages", mid, {"category": cat, "analysis_json": json.dumps(a), "analysis_source": "manual"})
    if d.get("route") and m["contact_id"]:
        contact = db.q1("SELECT * FROM contacts WHERE id=?", (m["contact_id"],))
        if m["enrollment_id"]:
            e = db.q1("SELECT * FROM enrollments WHERE id=?", (m["enrollment_id"],))
            if e and e["state"] in ("replied", "stopped", "completed") and e["phase"] == "done":
                db.update("enrollments", e["id"], {"state": "active"})
                e = db.q1("SELECT * FROM enrollments WHERE id=?", (e["id"],))
            if e and e["state"] in ("active", "paused"):
                engine.on_reply(e, m, a)
        worker.wake()
    return ok()


@app.post("/api/messages/<int:mid>/delete-conversation")
def delete_conversation(mid):
    """Removes a conversation from the app (Gmail is not touched). Sent emails stay as history."""
    d = body()
    m = db.q1("SELECT * FROM messages WHERE id=?", (mid,))
    if not m:
        abort(404)
    cid = m["contact_id"]
    if cid:
        ids = [r["id"] for r in db.q("SELECT id FROM messages WHERE contact_id=? AND direction='in'", (cid,))]
    else:
        ids = [r["id"] for r in db.q("SELECT id FROM messages WHERE direction='in' AND contact_id IS NULL AND from_addr=?",
                                     (m["from_addr"],))]
    resumed = 0
    with db.transaction() as cx:
        for i in ids:
            cx.execute("UPDATE events SET message_row_id=NULL WHERE message_row_id=?", (i,))
            cx.execute("UPDATE messages SET replied_to_id=NULL WHERE replied_to_id=?", (i,))
            cx.execute("DELETE FROM messages WHERE id=?", (i,))
    if cid and d.get("resume"):
        for e in db.q("SELECT * FROM enrollments WHERE contact_id=? AND state IN ('replied','stopped') "
                      "AND stopped_reason LIKE 'reply%'", (cid,)):
            db.update("enrollments", e["id"], {"state": "active", "phase": "waiting", "next_action_at": db.now_iso(),
                                               "stopped_reason": "", "updated_at": db.now_iso()})
            resumed += 1
    if cid:
        c = db.q1("SELECT * FROM contacts WHERE id=?", (cid,))
        sent = db.scalar("SELECT COUNT(*) FROM messages WHERE contact_id=? AND direction='out'", (cid,))
        used = db.scalar("SELECT COUNT(*) FROM enrollments WHERE contact_id=?", (cid,)) or \
            db.scalar("SELECT COUNT(*) FROM batch_members WHERE contact_id=?", (cid,))
        if c and "inbound" in (c["tags"] or "") and not sent and not used:
            co_id = c["company_id"]
            db.execute("UPDATE events SET contact_id=NULL WHERE contact_id=?", (cid,))
            db.execute("DELETE FROM contacts WHERE id=?", (cid,))
            if co_id:
                db.execute("DELETE FROM companies WHERE id=? AND plan='inbound' AND id NOT IN "
                           "(SELECT company_id FROM contacts WHERE company_id IS NOT NULL)", (co_id,))
        elif c and c["status"] in ("replied", "interested", "not_interested"):
            db.update("contacts", cid, {"status": "active" if sent else "new", "last_reply_at": None})
    if d.get("ignore"):
        lst = db.get_setting("ignored_senders") or []
        if m["from_addr"] and m["from_addr"] not in lst:
            lst.append(m["from_addr"].lower())
            db.set_setting("ignored_senders", lst)
    worker.wake()
    return ok(deleted=len(ids), resumed=resumed)


@app.get("/api/ignored-senders")
def ignored_senders():
    return jsonify(db.get_setting("ignored_senders") or [])


@app.post("/api/ignored-senders")
def set_ignored_senders():
    db.set_setting("ignored_senders", [x.strip().lower() for x in body().get("list") or [] if x.strip()])
    return ok()


@app.post("/api/messages/<int:mid>/find-price-list")
def find_price_list(mid):
    from . import crawler
    crawler.start(mid)
    return ok()


@app.post("/api/messages/<int:mid>/analyze")
def reanalyze(mid):
    if not analyzer.ai_available():
        raise ValueError("Add an Anthropic API key in Settings first.")
    engine.analyze_and_route(mid, route=False, use_ai=True)
    m = db.q1("SELECT analysis_source, analysis_json FROM messages WHERE id=?", (mid,))
    if m["analysis_source"] != "claude":
        raise RuntimeError("Claude could not analyse this email: " + (json.loads(m["analysis_json"] or "{}").get("next_action") or ""))
    return ok()


@app.post("/api/messages/<int:mid>/draft")
def draft(mid):
    m = db.q1("SELECT * FROM messages WHERE id=?", (mid,))
    if not analyzer.ai_available():
        raise ValueError("Add an Anthropic API key in Settings to draft replies with Claude.")
    text = analyzer.draft_reply(m, body().get("instructions") or "")
    acc = db.q1("SELECT * FROM accounts WHERE id=?", (m["account_id"],)) if m["account_id"] else None
    if acc and (acc.get("signature") or "").strip() and body().get("signature", True):
        text = text.rstrip() + "\n\n" + acc["signature"].strip()
    return ok(body=text)


@app.post("/api/send")
def send():
    d = body()
    if not d.get("to") or not d.get("subject") or not d.get("body"):
        raise ValueError("To, subject and message are required.")
    qid = engine.compose(d.get("account_id"), d["to"], d["subject"], d["body"], d.get("reply_to_message_id"),
                         d.get("send_at"), d.get("contact_id"))
    return ok(queue_id=qid)


# ---------------------------------------------------------------- stats & reports

@app.get("/api/stats")
def get_stats():
    pid = request.args.get("pipeline_id") or None
    days = int(request.args.get("days") or 30)
    return jsonify({"summary": stats.summary(pid, days), "daily": stats.daily(pid, min(days, 120)),
                    "send_time": stats.by_send_time(pid, max(days, 30)),
                    "steps": stats.by_step(int(pid)) if pid else [],
                    "pipelines": stats.pipeline_overview(), "leaders": stats.leaderboard(days)})


@app.get("/api/export/<kind>.csv")
def export(kind):
    if kind == "contacts":
        rows = db.q("SELECT c.email, c.first_name, c.last_name, c.title, co.name AS company, co.domain, c.status, c.tags, "
                    "c.last_contacted_at, c.last_reply_at, c.notes FROM contacts c LEFT JOIN companies co ON co.id=c.company_id")
    elif kind == "messages":
        rows = db.q("SELECT m.ts, m.direction, m.from_addr, m.to_addr, co.name AS company, p.name AS pipeline, m.node_id AS step, "
                    "m.subject, m.category, m.sentiment, m.interest_score, m.summary, m.is_test FROM messages m "
                    "LEFT JOIN companies co ON co.id=m.company_id LEFT JOIN pipelines p ON p.id=m.pipeline_id ORDER BY m.ts")
    elif kind == "companies":
        rows = db.q(f"SELECT co.domain, co.website, co.source, co.plan, {COMPANY_STATUS_SQL} AS status, co.form_url, "
                    f"co.form_status, co.brands, co.category, co.notes, "
                    f"(SELECT c.email FROM contacts c WHERE c.company_id=co.id ORDER BY c.id LIMIT 1) AS email "
                    f"FROM companies co WHERE COALESCE(co.plan,'')<>'' ORDER BY co.source, co.domain")
    elif kind == "replies":
        rows = db.q("SELECT m.ts AS received, m.from_addr, co.name AS company, p.name AS pipeline, m.node_id AS step, "
                    "m.category, m.sentiment, m.interest_score, m.summary, "
                    "(SELECT o.ts FROM messages o WHERE o.id=m.replied_to_id) AS our_email_sent, "
                    "ROUND((julianday(m.ts) - julianday((SELECT o.ts FROM messages o WHERE o.id=m.replied_to_id)))*24, 1) "
                    "AS hours_to_reply FROM messages m LEFT JOIN companies co ON co.id=m.company_id "
                    "LEFT JOIN pipelines p ON p.id=m.pipeline_id WHERE m.direction='in' ORDER BY m.ts")
    else:
        abort(404)
    buf = io.StringIO()
    if rows:
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    return Response(buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename=emailmanager-{kind}.csv"})


# ---------------------------------------------------------------- notifications

@app.get("/api/notifications")
def notifications():
    return jsonify(db.q("SELECT * FROM notifications ORDER BY id DESC LIMIT 60"))


@app.post("/api/notifications/read")
def notifications_read():
    db.execute("UPDATE notifications SET read=1 WHERE read=0")
    return ok()


@app.post("/api/notifications/test")
def notifications_test():
    from .notifier import notify
    notify("summary", "EmailManager test", "Notifications work. You'll get alerts like this for replies.", "#/", "info")
    return ok()


@app.get("/api/template-vars")
def template_vars():
    custom = set()
    for r in db.q("SELECT custom_json FROM contacts WHERE custom_json<>'{}' LIMIT 500"):
        custom |= set(json.loads(r["custom_json"]).keys())
    base = ["first_name", "last_name", "full_name", "email", "title", "company", "category", "website", "domain",
            "sender_name", "sender_first_name", "day_of_week", "original_subject"]
    return jsonify({"base": base, "custom": sorted(custom)})


# ---------------------------------------------------------------- simple sequences

def _seq_summary(p):
    d = sequences.to_api(p)
    d["errors"] = sequences.validate(d["steps"])
    d["people"] = db.scalar("SELECT COUNT(*) FROM enrollments WHERE pipeline_id=?", (p["id"],)) or 0
    d["running"] = db.scalar("SELECT COUNT(*) FROM enrollments WHERE pipeline_id=? AND state='active'", (p["id"],)) or 0
    s = stats.summary(p["id"], days=3650)
    d.update(sent=s["sent"], replied=s["replied"], reply_rate=s["reply_rate"], interested=s["interested"])
    return d


@app.get("/api/sequences")
def list_sequences():
    return jsonify([_seq_summary(p) for p in db.q("SELECT * FROM pipelines WHERE status<>'archived' ORDER BY id DESC")])


@app.get("/api/sequences/<int:pid>")
def get_sequence(pid):
    p = db.q1("SELECT * FROM pipelines WHERE id=?", (pid,))
    if not p:
        abort(404)
    return jsonify(_seq_summary(p))


@app.post("/api/sequences")
def create_sequence():
    d = body()
    steps = d.get("steps") or [{"subject": "", "body": "Hi {{first_name|there}},\n\n", "wait_days": 3}]
    pid, errs = sequences.save(None, d.get("name") or "New sequence", steps, d.get("settings") or {})
    return ok(id=pid, errors=errs)


@app.put("/api/sequences/<int:pid>")
def update_sequence(pid):
    d = body()
    _, errs = sequences.save(pid, d.get("name"), d.get("steps") or [], d.get("settings") or {})
    worker.wake()
    return ok(errors=errs)


@app.post("/api/sequences/<int:pid>/duplicate")
def duplicate_sequence(pid):
    p = sequences.to_api(db.q1("SELECT * FROM pipelines WHERE id=?", (pid,)))
    steps = [{k: v for k, v in s.items() if k != "id"} for s in p["steps"]]
    nid, _ = sequences.save(None, p["name"] + " (copy)", steps, p["settings"])
    return ok(id=nid)


@app.delete("/api/sequences/<int:pid>")
def delete_sequence(pid):
    return delete_pipeline(pid)


# ---------------------------------------------------------------- batches

def _batch_counts(bid):
    states = {r["state"]: r["n"] for r in db.q("SELECT state, COUNT(*) AS n FROM enrollments WHERE batch_id=? GROUP BY state", (bid,))}
    total = sum(states.values())
    sent = db.scalar("SELECT COUNT(*) FROM messages m JOIN enrollments e ON e.id=m.enrollment_id "
                     "WHERE e.batch_id=? AND m.direction='out'", (bid,)) or 0
    replied = db.scalar("SELECT COUNT(DISTINCT e.contact_id) FROM enrollments e JOIN messages m ON m.contact_id=e.contact_id "
                        "WHERE e.batch_id=? AND m.direction='in' AND m.ts>=e.created_at AND m.is_bounce=0 "
                        "AND COALESCE(m.category,'') NOT IN ('auto_reply','out_of_office','bounce')", (bid,)) or 0
    interested = db.scalar("SELECT COUNT(DISTINCT e.contact_id) FROM enrollments e JOIN messages m ON m.contact_id=e.contact_id "
                           "WHERE e.batch_id=? AND m.direction='in' AND m.ts>=e.created_at "
                           "AND m.category IN ('interested','question','referral')", (bid,)) or 0
    contacted = db.scalar("SELECT COUNT(DISTINCT e.id) FROM enrollments e JOIN messages m ON m.enrollment_id=e.id "
                          "WHERE e.batch_id=? AND m.direction='out'", (bid,)) or 0
    nxt = db.scalar("SELECT MIN(q.scheduled_at) FROM queue q JOIN enrollments e ON e.id=q.enrollment_id "
                    "WHERE e.batch_id=? AND q.status='pending'", (bid,))
    return {"total": total, "running": states.get("active", 0), "paused": states.get("paused", 0),
            "finished": states.get("completed", 0), "bounced": states.get("bounced", 0) + states.get("unsubscribed", 0),
            "sent": sent, "contacted": contacted, "replied": replied, "positive": interested,
            "reply_rate": round(100 * replied / contacted, 1) if contacted else 0, "next_at": nxt}


def _batch_row(b):
    acc = db.q1("SELECT email, from_name FROM accounts WHERE id=?", (b["account_id"],)) or {}
    p = db.q1("SELECT name FROM pipelines WHERE id=?", (b["pipeline_id"],)) or {}
    row = {**b, "from_email": acc.get("email", ""), "from_name": acc.get("from_name", ""), "sequence": p.get("name", ""),
           **_batch_counts(b["id"])}
    try:
        row["settings"] = json.loads(b.get("settings") or "{}")
    except Exception:
        row["settings"] = {}
    if b["status"] == "draft":
        row["total"] = db.scalar("SELECT COUNT(*) FROM batch_members WHERE batch_id=?", (b["id"],)) or 0
    # finished = stopped by hand, or nobody left in progress (everyone replied, finished or bounced)
    row["is_done"] = b["status"] == "stopped" or (b["status"] != "draft" and not row["running"] and not row["paused"])
    if row["is_done"]:
        row["done_at"] = db.scalar("SELECT MAX(e.updated_at) FROM enrollments e WHERE e.batch_id=?", (b["id"],)) or b["created_at"]
    return row


@app.get("/api/batches")
def list_batches():
    rows = db.q("SELECT * FROM batches ORDER BY CASE status WHEN 'draft' THEN 1 ELSE 0 END, "
                "CASE WHEN status='draft' THEN priority ELSE 0 END, CASE WHEN status='draft' THEN id ELSE -id END")
    return jsonify([_batch_row(b) for b in rows])


@app.post("/api/batches/parse")
def parse_batch():
    recipients, errors = sequences.parse_recipients(body().get("text") or "")
    return ok(recipients=recipients, errors=errors)


@app.post("/api/batches/preview")
def preview_batch():
    d = body()
    if not d.get("pipeline_id"):
        raise ValueError("Choose a sequence.")
    return ok(emails=sequences.render_preview(int(d["pipeline_id"]), d.get("recipient") or {}, d.get("account_id")))


def _add_people(bid, b, recipients, start_at=None):
    cids = []
    for r in recipients:
        cid = engine.upsert_contact({"email": r["email"], "first_name": r.get("first_name", ""),
                                     "last_name": r.get("last_name", ""), "company": r.get("company", ""),
                                     "website": r.get("website", ""), "title": r.get("title", ""),
                                     "category": r.get("category", ""), "custom": r.get("custom") or {},
                                     "tags": f"batch-{bid}"})
        cids.append(cid)
        db.execute("INSERT OR IGNORE INTO batch_members (batch_id, contact_id, added_at) VALUES (?,?,?)",
                   (bid, cid, db.now_iso()))
    return engine.enroll(b["pipeline_id"], cids, start_at, 0, b["account_id"], batch_id=bid)


@app.post("/api/batches")
def create_batch():
    d = body()
    pid, acc_id = d.get("pipeline_id"), d.get("account_id")
    recipients = d.get("recipients") or []
    if not pid:
        raise ValueError("Choose a sequence.")
    if not acc_id:
        raise ValueError("Choose which mailbox to send from.")
    if not recipients:
        raise ValueError("Add at least one email address.")
    p = db.q1("SELECT * FROM pipelines WHERE id=?", (pid,))
    errs = sequences.validate(sequences.to_api(p)["steps"])
    if errs:
        raise ValueError("Fix the sequence first: " + " ".join(errs))
    db.update("pipelines", pid, {"status": "active"})
    name = (d.get("name") or "").strip() or f"{p['name']} · {db.utcnow().strftime('%b %d')}"
    bid = db.insert("batches", {"name": name, "pipeline_id": pid, "account_id": acc_id, "status": "running",
                                "start_at": d.get("start_at") or db.now_iso(), "created_at": db.now_iso()})
    res = _add_people(bid, {"pipeline_id": pid, "account_id": acc_id}, recipients, d.get("start_at"))
    worker.wake()
    return ok(id=bid, **res)


@app.get("/api/batches/<int:bid>")
def get_batch(bid):
    b = db.q1("SELECT * FROM batches WHERE id=?", (bid,))
    if not b:
        abort(404)
    out = _batch_row(b)
    if b["status"] == "draft":
        out["people"] = db.q("SELECT c.id AS contact_id, c.email, c.first_name, c.last_name, c.status AS contact_status, "
                             "c.custom_json, co.domain, co.website, co.brands, co.category, co.crm_note "
                             "FROM batch_members m JOIN contacts c ON c.id=m.contact_id "
                             "LEFT JOIN companies co ON co.id=c.company_id WHERE m.batch_id=? ORDER BY co.domain", (bid,))
        for r in out["people"]:
            r["skip"] = engine.already_reached(r | {"id": r["contact_id"], "status": r["contact_status"],
                                                    "company_id": db.scalar("SELECT company_id FROM contacts WHERE id=?",
                                                                            (r["contact_id"],))})
        return jsonify(out)
    p = db.q1("SELECT * FROM pipelines WHERE id=?", (b["pipeline_id"],))
    steps = sequences.to_api(p)["steps"] if p else []
    rows = db.q("SELECT e.*, c.email, c.first_name, c.last_name, co.name AS company FROM enrollments e "
                "JOIN contacts c ON c.id=e.contact_id LEFT JOIN companies co ON co.id=c.company_id "
                "WHERE e.batch_id=? ORDER BY e.id", (bid,))
    now = db.now_iso()
    out["people"] = [{"id": e["id"], "contact_id": e["contact_id"], "email": e["email"], "first_name": e["first_name"],
                      "last_name": e["last_name"], "company": e["company"], "state": e["state"],
                      **sequences.recipient_status(e, steps, now)} for e in rows]
    return jsonify(out)


@app.post("/api/batches/<int:bid>/start")
def start_batch(bid):
    d = body()
    b = db.q1("SELECT * FROM batches WHERE id=?", (bid,))
    if not b or b["status"] != "draft":
        raise ValueError("This batch has already started.")
    pid, acc_id = d.get("pipeline_id"), d.get("account_id")
    if not pid:
        raise ValueError("Choose an email style.")
    if not acc_id:
        raise ValueError("Choose which mailbox to send from.")
    p = db.q1("SELECT * FROM pipelines WHERE id=?", (pid,))
    errs = sequences.validate(sequences.to_api(p)["steps"])
    if errs:
        raise ValueError("Fix the email style first: " + " ".join(errs))
    db.update("pipelines", pid, {"status": "active"})
    db.update("batches", bid, {"pipeline_id": pid, "account_id": acc_id, "status": "running",
                               "start_at": d.get("start_at") or db.now_iso(),
                               "settings": json.dumps(_clean_batch_settings(d.get("settings") or {}))})
    cids = [r["contact_id"] for r in db.q("SELECT contact_id FROM batch_members WHERE batch_id=?", (bid,))]
    res = engine.enroll(pid, cids, d.get("start_at"), 0, acc_id, batch_id=bid)
    worker.wake()
    return ok(**res)


def _clean_batch_settings(st):
    out = {}
    days = [int(x) for x in st.get("days") or [] if str(x).isdigit() and 0 <= int(x) <= 6]
    if days:
        out["days"] = sorted(set(days))
    for k in ("window_start", "window_end"):
        v = str(st.get(k) or "")
        if re.fullmatch(r"\d\d:\d\d", v):
            out[k] = v
    if out.get("window_start") and out.get("window_end") and out["window_start"] >= out["window_end"]:
        raise ValueError("The send window must end after it starts.")
    if st.get("daily_limit") not in (None, "", 0, "0"):
        out["daily_limit"] = max(1, int(st["daily_limit"]))
    return out


@app.post("/api/batches/<int:bid>/settings")
def batch_settings(bid):
    db.update("batches", bid, {"settings": json.dumps(_clean_batch_settings(body()))})
    b = db.q1("SELECT account_id FROM batches WHERE id=?", (bid,))
    if b and b["account_id"]:
        engine.space_out(b["account_id"])
    worker.wake()
    return ok()


@app.post("/api/batches/<int:bid>/remove")
def remove_from_batch(bid):
    for cid in body().get("contact_ids") or []:
        db.execute("DELETE FROM batch_members WHERE batch_id=? AND contact_id=? AND batch_id IN "
                   "(SELECT id FROM batches WHERE status='draft')", (bid, int(cid)))
    return ok()


# ---------------------------------------------------------------- companies & progress

COMPANY_STATUS_SQL = """
CASE
  WHEN co.plan='contacted_before' THEN 'contacted_before'
  WHEN co.plan='form' THEN CASE co.form_status WHEN 'submitted' THEN 'form_submitted' WHEN 'not_possible' THEN 'form_not_possible' ELSE 'form_todo' END
  WHEN co.plan IN ('duplicate','no_contact') THEN co.plan
  WHEN EXISTS (SELECT 1 FROM messages m WHERE m.company_id=co.id AND m.direction='in' AND m.category='interested') THEN 'interested'
  WHEN EXISTS (SELECT 1 FROM messages m WHERE m.company_id=co.id AND m.direction='in' AND m.is_bounce=1) THEN 'bounced'
  WHEN EXISTS (SELECT 1 FROM messages m WHERE m.company_id=co.id AND m.direction='in'
               AND COALESCE(m.category,'') NOT IN ('auto_reply','out_of_office','bounce')) THEN 'replied'
  WHEN EXISTS (SELECT 1 FROM enrollments e JOIN contacts c ON c.id=e.contact_id WHERE c.company_id=co.id AND e.state='completed') THEN 'no_reply'
  WHEN EXISTS (SELECT 1 FROM messages m WHERE m.company_id=co.id AND m.direction='out' AND m.is_test=0) THEN 'emailing'
  WHEN EXISTS (SELECT 1 FROM enrollments e JOIN contacts c ON c.id=e.contact_id WHERE c.company_id=co.id) THEN 'scheduled'
  WHEN EXISTS (SELECT 1 FROM batch_members bm JOIN contacts c ON c.id=bm.contact_id WHERE c.company_id=co.id) THEN 'ready'
  ELSE 'not_in_batch'
END"""


@app.get("/api/companies/overview")
def companies_overview():
    rows = db.q(f"SELECT co.source, {COMPANY_STATUS_SQL} AS st, COUNT(*) AS n FROM companies co "
                f"WHERE COALESCE(co.plan,'')<>'' GROUP BY 1, 2")
    by = {}
    for r in rows:
        by.setdefault(r["source"] or "Added by hand", {})[r["st"]] = r["n"]
    return jsonify(by)


@app.get("/api/companies/list")
def companies_list():
    where, args = ["COALESCE(co.plan,'')<>''"], []
    if request.args.get("source"):
        where.append("co.source=?")
        args.append(request.args["source"])
    if request.args.get("q"):
        like = f"%{request.args['q'].lower()}%"
        where.append("(co.domain LIKE ? OR EXISTS (SELECT 1 FROM contacts c WHERE c.company_id=co.id AND c.email LIKE ?))")
        args += [like, like]
    fs = request.args.get("fs")
    if fs is not None and fs in ("todo", "submitted", "not_possible"):
        where.append("co.plan='form'")
        where.append("COALESCE(co.form_status,'')=?")
        args.append("" if fs == "todo" else fs)
    st = request.args.get("status") or ""
    groups = {"email": "('ready','not_in_batch','scheduled','emailing','no_reply')",
              "replied": "('replied','interested','bounced')", "forms": "('form_todo','form_submitted','form_not_possible')",
              "before": "('contacted_before')", "skipped": "('duplicate','no_contact')"}
    w = " AND ".join(where)
    inner = (f"SELECT co.*, {COMPANY_STATUS_SQL} AS st, "
             f"(SELECT c.email FROM contacts c WHERE c.company_id=co.id ORDER BY c.id LIMIT 1) AS email, "
             f"(SELECT b.name FROM batch_members bm JOIN contacts c ON c.id=bm.contact_id JOIN batches b ON b.id=bm.batch_id "
             f" WHERE c.company_id=co.id LIMIT 1) AS batch, "
             f"(SELECT MAX(m.ts) FROM messages m WHERE m.company_id=co.id) AS last_ts, "
             f"(SELECT m.id FROM messages m WHERE m.company_id=co.id AND m.direction='in' ORDER BY m.ts DESC LIMIT 1) AS reply_id "
             f"FROM companies co WHERE {w}")
    outer = f"SELECT * FROM ({inner}) t" + (f" WHERE st IN {groups[st]}" if st in groups else "")
    total = db.scalar(f"SELECT COUNT(*) FROM ({outer})", tuple(args))
    page = max(0, int(request.args.get("page") or 0))
    order = ("CASE source WHEN 'Paid catalogs' THEN 1 WHEN 'Astra brand suppliers' THEN 2 WHEN 'SpyFu competitors' THEN 3 "
             "ELSE 4 END, domain") if fs is not None else \
        "CASE WHEN last_ts IS NULL THEN 1 ELSE 0 END, last_ts DESC, domain"
    rows = db.q(outer + f" ORDER BY {order} LIMIT 100 OFFSET ?", tuple(args) + (page * 100,))
    sources = [r["source"] for r in db.q("SELECT DISTINCT source FROM companies WHERE COALESCE(source,'')<>'' ORDER BY source")]
    return jsonify({"total": total, "rows": rows, "sources": sources})


@app.post("/api/companies/<int:cid>/note")
def company_note(cid):
    db.update("companies", cid, {"notes": (body().get("note") or "").strip()[:2000], "updated_at": db.now_iso()})
    return ok()


# ---------------------------------------------------------------- price lists (attachments & links in replies)

@app.get("/api/price-lists")
def price_lists():
    st = request.args.get("status") or "new"
    where = "" if st == "all" else "WHERE pl.status=?"
    keepa.init()
    rows = db.q(f"SELECT pl.*, m.subject, m.from_addr, m.from_name, co.domain, co.name AS company, "
                f"(SELECT a.id FROM analyses a WHERE a.price_list_id=pl.id ORDER BY a.id DESC LIMIT 1) AS analysis_id "
                f"FROM price_lists pl LEFT JOIN messages m ON m.id=pl.message_row_id "
                f"LEFT JOIN companies co ON co.id=pl.company_id {where} ORDER BY pl.received_at DESC LIMIT 500",
                () if st == "all" else (st,))
    counts = {r["status"]: r["n"] for r in db.q("SELECT status, COUNT(*) AS n FROM price_lists GROUP BY status")}
    return jsonify({"rows": rows, "counts": counts})


@app.get("/api/price-lists/<int:pid>/file")
def price_list_file(pid):
    r = db.q1("SELECT * FROM price_lists WHERE id=? AND kind='file'", (pid,))
    if not r:
        abort(404)
    root = (db.DATA_DIR / "attachments").resolve()
    path = __import__("pathlib").Path(r["path"]).resolve()
    if root not in path.parents or not path.exists():
        abort(404)
    inline = r["content_type"] in ("application/pdf",) or r["content_type"].startswith("image/") or \
        r["content_type"].startswith("text/")
    return send_file(path, mimetype=r["content_type"] or None, as_attachment=not inline, download_name=r["filename"])


@app.get("/api/price-lists/<int:pid>/preview")
def price_list_preview(pid):
    """What the browser can show: PDFs/images inline, spreadsheets and Word files as tables/text, ZIPs as a file list."""
    import csv as _csv
    import zipfile
    from pathlib import Path
    r = db.q1("SELECT * FROM price_lists WHERE id=?", (pid,))
    if not r:
        abort(404)
    if r["kind"] == "link":
        return jsonify({"type": "link", "url": r["url"]})
    path = Path(r["path"]).resolve()
    if (db.DATA_DIR / "attachments").resolve() not in path.parents or not path.exists():
        abort(404)
    name, ctype = r["filename"].lower(), r["content_type"] or ""
    url = f"/api/price-lists/{pid}/file"
    try:
        if ctype == "application/pdf" or name.endswith(".pdf"):
            return jsonify({"type": "pdf", "url": url})
        if ctype.startswith("image/"):
            return jsonify({"type": "image", "url": url})
        if name.endswith((".xlsx", ".xlsm")):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheets = []
            for ws in wb.worksheets[:8]:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    if len(rows) >= 500:
                        break
                    vals = ["" if v is None else str(v) for v in row[:40]]
                    if any(vals):
                        rows.append(vals)
                sheets.append({"name": ws.title, "rows": rows})
            return jsonify({"type": "table", "sheets": sheets})
        if name.endswith((".csv", ".txt")) or ctype.startswith("text/"):
            text = path.read_bytes()[:2_000_000].decode("utf-8", "ignore")
            if name.endswith(".csv") or ctype == "text/csv":
                try:
                    dialect = _csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
                except Exception:
                    dialect = _csv.excel
                rows = [r_[:40] for r_ in list(_csv.reader(text.splitlines(), dialect))[:500]]
                return jsonify({"type": "table", "sheets": [{"name": r["filename"], "rows": rows}]})
            return jsonify({"type": "text", "text": text[:200_000]})
        if name.endswith(".docx"):
            import docx
            d = docx.Document(str(path))
            parts = [p_.text for p_ in d.paragraphs if p_.text.strip()]
            sheets = [{"name": f"Table {i + 1}", "rows": [[c.text for c in row.cells][:40] for row in t.rows][:500]}
                      for i, t in enumerate(d.tables[:8])]
            return jsonify({"type": "doc", "text": "\n".join(parts)[:200_000], "sheets": sheets})
        if name.endswith(".zip") or ctype == "application/zip":
            with zipfile.ZipFile(path) as z:
                return jsonify({"type": "zip", "files": [{"name": i.filename, "size": i.file_size} for i in z.infolist()[:300]]})
    except Exception as ex:
        return jsonify({"type": "none", "url": url, "error": f"Can't show this file here ({ex}). Download it instead."})
    return jsonify({"type": "none", "url": url, "error": "This file type can't be shown in the browser. Download it instead."})


@app.post("/api/price-lists/<int:pid>")
def update_price_list(pid):
    d = body()
    upd = {k: d[k] for k in ("status", "notes") if k in d}
    if upd.get("status") not in (None, "new", "reviewed", "not_price_list"):
        raise ValueError("bad status")
    db.update("price_lists", pid, upd)
    return ok()


# ---------------------------------------------------------------- Keepa analysis of price lists

@app.post("/api/price-lists/upload")
def upload_price_list():
    f = request.files.get("file")
    if not f or not f.filename:
        raise ValueError("Choose a file.")
    import re as _re
    folder = db.DATA_DIR / "attachments" / "uploads" / db.utcnow().strftime("%Y%m%d%H%M%S")
    folder.mkdir(parents=True, exist_ok=True)
    safe = _re.sub(r"[^\w.\- ()]+", "_", f.filename)[:120]
    path = folder / safe
    f.save(path)
    pid = db.insert("price_lists", {"kind": "file", "filename": f.filename, "path": str(path), "size": path.stat().st_size,
                                    "content_type": f.mimetype or "", "received_at": db.now_iso(), "status": "new",
                                    "notes": (request.form.get("supplier") or "").strip()})
    return ok(id=pid)


@app.post("/api/price-lists/<int:pid>/analyse")
def analyse_price_list(pid):
    return ok(id=keepa.start(pid))


@app.get("/api/analyses/<int:aid>")
def get_analysis(aid):
    keepa.init()
    a = db.q1("SELECT * FROM analyses WHERE id=?", (aid,))
    if not a:
        abort(404)
    a["price_list"] = db.q1("SELECT pl.id, pl.filename, pl.notes, pl.message_row_id, co.domain, m.from_addr "
                            "FROM price_lists pl LEFT JOIN companies co ON co.id=pl.company_id "
                            "LEFT JOIN messages m ON m.id=pl.message_row_id WHERE pl.id=?", (a["price_list_id"],))
    a["items"] = db.q("SELECT * FROM analysis_items WHERE analysis_id=? ORDER BY row_no", (aid,))
    a["min_roi"] = db.get_setting("min_roi") or 30
    a["keepa_key_set"] = secrets.has_secret("keepa_api_key")
    return jsonify(a)


@app.post("/api/analyses/<int:aid>/keepa")
def analysis_rerun(aid):
    if not secrets.has_secret("keepa_api_key"):
        raise ValueError("Add your Keepa API key in Settings first.")
    keepa.token_status(secrets.get_secret("keepa_api_key"))  # fail fast with Keepa's own reason (plan, tokens)
    keepa.rerun_keepa(aid, by_name=bool(body().get("by_name")))
    return ok()


@app.post("/api/keepa/test")
def keepa_test():
    if not secrets.has_secret("keepa_api_key"):
        raise ValueError("No Keepa key saved yet.")
    return ok(**keepa.token_status(secrets.get_secret("keepa_api_key")))


@app.post("/api/analyses/<int:aid>/by-name")
def analysis_by_name(aid):
    if not secrets.has_secret("keepa_api_key"):
        raise ValueError("Add your Keepa API key in Settings first.")
    keepa.start_by_name(aid)
    return ok()


@app.get("/api/analyses/<int:aid>/export.csv")
def analysis_export(aid):
    rows = db.q("SELECT row_no, code AS upc_ean, sku, name, brand, cost, case_pack, msrp, moq, asin, amazon_title, "
                "price AS amazon_price, avg90 AS avg_price_90d, amazon_sells, sales_rank, monthly_sold, offers, "
                "offers_fba, referral_pct, fba_fee, profit, roi AS roi_pct, margin AS margin_pct, status "
                "FROM analysis_items WHERE analysis_id=? ORDER BY roi DESC", (aid,))
    buf = io.StringIO()
    if rows:
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    return Response(buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename=keepa-analysis-{aid}.csv"})


@app.get("/api/forms/counts")
def forms_counts():
    rows = db.q("SELECT COALESCE(form_status,'') AS fs, COUNT(*) AS n FROM companies WHERE plan='form' GROUP BY 1")
    return jsonify({("todo" if r["fs"] == "" else r["fs"]): r["n"] for r in rows})


@app.post("/api/companies/<int:cid>/form")
def company_form(cid):
    st = body().get("status") or ""
    if st not in ("", "submitted", "not_possible"):
        raise ValueError("bad status")
    db.update("companies", cid, {"form_status": st, "updated_at": db.now_iso()})
    return ok()


@app.post("/api/batches/<int:bid>/add")
def add_to_batch(bid):
    b = db.q1("SELECT * FROM batches WHERE id=?", (bid,))
    res = _add_people(bid, b, body().get("recipients") or [])
    worker.wake()
    return ok(**res)


@app.post("/api/batches/<int:bid>/action")
def batch_action(bid):
    a = body().get("action")
    rows = db.q("SELECT * FROM enrollments WHERE batch_id=?", (bid,))
    if a == "pause":
        for e in rows:
            if e["state"] == "active":
                db.update("enrollments", e["id"], {"state": "paused", "updated_at": db.now_iso()})
                db.execute("UPDATE queue SET status='needs_approval', approval_reason='batch paused' "
                           "WHERE enrollment_id=? AND status='pending'", (e["id"],))
        db.update("batches", bid, {"status": "paused"})
    elif a == "resume":
        for e in rows:
            if e["state"] == "paused":
                db.update("enrollments", e["id"], {"state": "active", "updated_at": db.now_iso()})
                db.execute("UPDATE queue SET status='pending', approval_reason='' WHERE enrollment_id=? "
                           "AND status='needs_approval' AND approval_reason IN ('batch paused','enrollment paused')", (e["id"],))
        db.update("batches", bid, {"status": "running"})
        for acc in db.q("SELECT id FROM accounts"):
            engine.space_out(acc["id"])
    elif a == "stop":
        for e in rows:
            if e["state"] in ("active", "paused"):
                engine.finish(e, "stop", "batch stopped by me")
        db.update("batches", bid, {"status": "stopped"})
    else:
        raise ValueError("unknown action")
    worker.wake()
    return ok()


def create_app():
    db.init()
    sequences.seed_examples()
    return app
