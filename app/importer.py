"""Import supplier spreadsheets into EmailManager.

    .venv/bin/python -m app.importer ../ResultsCompanies [--dry-run]

Every company gets exactly one plan:
  form              has a wholesale form on its website                                       -> forms xlsx, filled by hand
  email             has an email and no form                                                  -> goes into a ready batch
  duplicate         same email or same email domain as a company above                         -> not stored
  no_contact        neither email nor form
The CRM table is another team's outreach history: it is shown as background ("Other team: ...") but never
used to skip anyone. Re-running is safe; --reset first removes draft batches and untouched imported data.
"""
import json
import re
import sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd

from . import db, engine

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
BAD_EMAIL = re.compile(r"\.(png|jpe?g|gif|webp|svg)$|example\.|sentry|wixpress|noreply|no-reply|donotreply", re.I)
BATCH_SIZE = 100
# a form that can't be done online -> email instead
NOT_ONLINE = {"Отправка обычной почтой", "Не открывается сайт", "Нет простой формы для регистрации"}

SOURCES = {  # key: (label shown in the app, priority, original file)
    "product": ("Product suppliers", 1, "Доп поставщики без wholesale-форм.xlsx"),
    "catalog": ("Paid catalogs", 2, "Поставщики из платных каталогов + 2500.xlsx"),
    "competitor": ("Competitors of product suppliers", 3, "Конкуренты доп поставщиков.xlsx"),
    "astra": ("Astra brand suppliers", 4, "CRM таблица.xlsx · Astra"),
    "spyfu": ("SpyFu competitors", 5, "CRM таблица.xlsx · SpyFu"),
    "suppliers": ("Suppliers (first list)", 6, "CRM таблица.xlsx · Поставщики"),
    "offline_form": ("Form not possible online", 7, "CRM таблица.xlsx"),
}
CRM_SEGMENTS = [(2, 1701, "suppliers"), (1703, 2264, "spyfu"), (2265, 4653, "astra"), (4654, 6466, "catalog")]


def domain_of(url):
    if not isinstance(url, str) or not url.strip():
        return ""
    u = re.sub(r"^https?://", "", url.strip().lower())
    return u.split("/")[0].split("?")[0].split("#")[0].removeprefix("www.").strip(".")


def first_email(v):
    for e in EMAIL_RE.findall(str(v or "")):
        if not BAD_EMAIL.search(e):
            return e.lower().strip(".")
    return ""


def clean(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("nan", "none") else s


def fmt_date(v):
    v = clean(v)
    if not v:
        return ""
    if re.fullmatch(r"\d{5}(\.0)?", v):  # Excel serial
        return (datetime(1899, 12, 30) + timedelta(days=int(float(v)))).strftime("%Y-%m-%d")
    m = re.match(r"(\d{4})-(\d\d)-(\d\d)", v)
    if m:
        return m.group(0)
    m = re.match(r"(\d\d)\.(\d\d)\.(\d{4})", v)
    return f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else v[:10]


def is_url(v):
    return clean(v).lower().startswith("http")


# ---------------------------------------------------------------- read

def read_folder(folder):
    folder = Path(folder)
    crm_x = pd.read_excel(folder / "CRM таблица.xlsx", sheet_name=["общий", "CRM_лог", "CRM_Coldy"], dtype=str)
    crm = crm_x["общий"]
    log = crm_x["CRM_лог"].dropna(how="all")
    coldy = crm_x["CRM_Coldy"].dropna(how="all")
    log_by_dom = {}
    for r in log.itertuples(index=False):
        d = domain_of(getattr(r, "site", ""))
        if d:
            log_by_dom[d] = clean(getattr(r, "summary", ""))
    coldy_emails = {clean(e).lower() for e in coldy["supplier_email"] if clean(e)}

    companies = {}  # domain -> record

    def rec(dom, website, source):
        r = companies.get(dom)
        if r is None:
            r = companies[dom] = {"domain": dom, "website": website, "sources": [], "emails": [], "form_url": "",
                                  "crm": {}, "crm_contacted": False, "brands": [], "category": "", "product": "",
                                  "price_note": "", "form_reason": "", "crm_row": None}
        if source not in r["sources"]:
            r["sources"].append(source)
        if website and not r["website"]:
            r["website"] = website
        return r

    for idx, row in crm.iterrows():
        rownum = idx + 2
        site = clean(row.get("Сайт поставщика"))
        dom = domain_of(site)
        if not dom:
            continue
        seg = next((k for a, b, k in CRM_SEGMENTS if a <= rownum <= b), "suppliers")
        r = rec(dom, site, seg)
        r["crm_row"] = r["crm_row"] or rownum
        em = first_email(row.get("Почта"))
        if em and em not in r["emails"]:
            r["emails"].append(em)
        if is_url(row.get("Ссылка на форму Wholesale")) and not r["form_url"]:
            r["form_url"] = clean(row.get("Ссылка на форму Wholesale"))
        sent_form = clean(row.get("Отправили форму")).lower() == "true"
        emailed = fmt_date(row.get("Дата отправки 1 письма"))
        answer = clean(row.get("Получен ответ (статус)"))
        c = r["crm"]
        if sent_form:
            c["form_sent"] = True
        if emailed:
            c["emailed"] = emailed
        if answer:
            c["answer"] = re.sub(r"\s+", " ", answer)[:300]
        for k_src, k in (("Исполнитель", "who"), ("Активное действие от нас", "todo"),
                         ("Дата последнего взаимодействия", "last"), ("Причина почему не отправили", "reason"),
                         ("Пароль (если форма с регистрацией)", "password")):
            v = clean(row.get(k_src))
            if v and k not in c:
                c[k] = fmt_date(v) if k == "last" else v
        if dom in log_by_dom:
            c["log"] = log_by_dom[dom][:300]
        if sent_form or emailed or answer or dom in log_by_dom or (em and em in coldy_emails):
            r["crm_contacted"] = True

    cat = pd.read_excel(folder / "Поставщики из платных каталогов + 2500.xlsx", sheet_name="Лист1", dtype=str)
    for _, row in cat.iterrows():
        site = clean(row.get("Сайт поставщика"))
        dom = domain_of(site)
        if not dom:
            continue
        r = rec(dom, site, "catalog")
        em = first_email(row.get("Почта"))
        if em and em not in r["emails"]:
            r["emails"].append(em)
        if is_url(row.get("Ссылка на форму Wholesale (если есть)")) and not r["form_url"]:
            r["form_url"] = clean(row.get("Ссылка на форму Wholesale (если есть)"))
        r["price_note"] = clean(row.get("Открытые оптовые цены"))

    dop = pd.read_excel(folder / "Доп поставщики без wholesale-форм.xlsx", dtype=str)
    for _, row in dop.iterrows():
        site = clean(row.get("Сайт поставщика"))
        dom = domain_of(site)
        if not dom:
            continue
        r = rec(dom, site, "product")
        em = first_email(row.get("Email для wholesale-партнёров"))
        if em and em not in r["emails"]:
            r["emails"].append(em)
        b = clean(row.get("Brand")) or clean(row.get("Manufacturer"))
        if b and b.lower() not in [x.lower() for x in r["brands"]]:
            r["brands"].append(b)
        r["category"] = r["category"] or clean(row.get("Categories: Root"))
        r["product"] = r["product"] or clean(row.get("Title"))

    kon = pd.read_excel(folder / "Конкуренты доп поставщиков.xlsx", dtype=str)
    for _, row in kon.iterrows():
        site = clean(row.get("Сайт поставщика"))
        dom = domain_of(site)
        if not dom:
            continue
        r = rec(dom, site if site.startswith("http") else f"https://{site}", "competitor")
        em = first_email(row.get("Почта"))
        if em and em not in r["emails"]:
            r["emails"].append(em)
    return companies


# ---------------------------------------------------------------- decide

def crm_note(r):
    """What the other team did with this supplier (background only)."""
    c = r["crm"]
    parts = []
    if c.get("emailed"):
        parts.append(f"emailed {c['emailed']}")
    if c.get("form_sent"):
        parts.append("form submitted")
    if c.get("answer"):
        parts.append(f"reply: {c['answer']}")
    if c.get("todo"):
        parts.append(f"to do: {c['todo']}")
    if c.get("log"):
        parts.append(f"log: {c['log']}")
    if c.get("reason"):
        parts.append(f"form not sent: {c['reason']}")
    if c.get("who"):
        parts.append(f"by {c['who']}")
    return ("Other team: " + " · ".join(parts)) if parts else ""


def plan(companies):
    """Assigns r['plan'], r['email'], r['source'] and returns ordered email candidates."""
    used_emails, used_email_domains = {}, {}

    def email_dom(e):
        d = e.split("@")[-1]
        return None if d in engine.FREEMAIL else d

    def claim(r, e):
        used_emails[e] = r["domain"]
        d = email_dom(e)
        if d:
            used_email_domains.setdefault(d, r["domain"])
        used_email_domains.setdefault(r["domain"], r["domain"])

    order = sorted(companies.values(), key=lambda r: min(SOURCES[s][1] for s in r["sources"]))
    # 1) forms first - they block their emails/domains from the email batches.
    # CRM history is NOT used to skip anyone: it is another team's outreach, Oleg has contacted nobody.
    for r in order:
        r["source"] = min(r["sources"], key=lambda s: SOURCES[s][1])
        r["email"] = r["emails"][0] if r["emails"] else ""
    for r in order:
        if r.get("plan"):
            continue
        reason = r["crm"].get("reason", "")
        if r["form_url"] and reason not in NOT_ONLINE:
            r["plan"] = "form"
            for e in r["emails"]:
                claim(r, e)
        elif r["form_url"] and reason in NOT_ONLINE:
            r["source"] = "offline_form"
    # 2) emails, best sources first
    for r in order:
        if r.get("plan"):
            continue
        if not r["emails"]:
            r["plan"] = "no_contact"
            continue
        e = next((x for x in r["emails"] if x not in used_emails), None)
        if e is None:
            r["plan"], r["dup_of"] = "duplicate", used_emails[r["emails"][0]]
            continue
        d = email_dom(e)
        other = used_email_domains.get(d) if d else None
        if (other and other != r["domain"]) or (r["domain"] in used_email_domains and used_email_domains[r["domain"]] != r["domain"]):
            r["plan"], r["dup_of"] = "duplicate", other or used_email_domains[r["domain"]]
            continue
        r["plan"], r["email"] = "email", e
        claim(r, e)
    return order


# ---------------------------------------------------------------- write to the app database

def batch_groups(order):
    groups = defaultdict(list)
    for r in order:
        if r["plan"] != "email":
            continue
        key = (r["source"], r["category"] if r["source"] == "product" else "")
        groups[key].append(r)
    out = []
    for (src, catg), rs in sorted(groups.items(), key=lambda kv: (SOURCES[kv[0][0]][1], -len(kv[1]))):
        label, prio, file = SOURCES[src]
        if src == "product" and len(rs) < 40:  # small categories go together
            catg = "Other categories"
        out.append((src, catg, rs))
    merged = defaultdict(list)
    for src, catg, rs in out:
        merged[(src, catg)].extend(rs)
    final = []
    for (src, catg), rs in merged.items():
        n = -(-len(rs) // BATCH_SIZE)
        size = -(-len(rs) // n)
        chunks = [rs[i:i + size] for i in range(0, len(rs), size)]
        for i, ch in enumerate(chunks, start=1):
            label = SOURCES[src][0] + (f" · {catg}" if catg else "") + (f" · {i}/{len(chunks)}" if len(chunks) > 1 else "")
            final.append((src, label, ch))
    return final


def apply(order):
    now = db.now_iso()
    stats = defaultdict(int)
    with db.transaction() as cx:
        for r in order:
            if r["plan"] == "duplicate":
                cx.execute("DELETE FROM companies WHERE domain=? AND id NOT IN (SELECT company_id FROM contacts "
                           "WHERE company_id IS NOT NULL)", (r["domain"],))
                stats["duplicate (deleted)"] += 1
                continue
            existing = cx.execute("SELECT id, form_status FROM companies WHERE domain=?", (r["domain"],)).fetchone()
            fields = {"website": r["website"], "source": SOURCES[r["source"]][0], "plan": r["plan"],
                      "form_url": r["form_url"], "crm_note": crm_note(r), "crm_row": r["crm_row"],
                      "brands": ", ".join(r["brands"][:3]), "category": r["category"], "updated_at": now}
            if existing:
                cid = existing[0]
                sets = ",".join(f"{k}=?" for k in fields)
                cx.execute(f"UPDATE companies SET {sets} WHERE id=?", tuple(fields.values()) + (cid,))
            else:
                cols = ["name", "domain", "notes", "created_at", *fields]
                vals = ["", r["domain"], "", now, *fields.values()]
                cid = cx.execute(f"INSERT INTO companies ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})", vals).lastrowid
            r["company_id"] = cid
            stats[r["plan"]] += 1
            # one contact per email
            for e in r["emails"]:
                custom = {}
                if r["brands"]:
                    custom["brands"] = " and ".join(r["brands"][:2]) if len(r["brands"]) <= 2 else ", ".join(r["brands"][:2])
                if r["category"]:
                    custom["category_name"] = r["category"]
                if r["product"]:
                    custom["product"] = r["product"][:120]
                status = "contacted_before" if r["plan"] == "contacted_before" else "new"
                row = cx.execute("SELECT id, status, company_id FROM contacts WHERE email=?", (e,)).fetchone()
                if row:
                    if row[1] == "new" and status == "contacted_before":
                        cx.execute("UPDATE contacts SET status=? WHERE id=?", (status, row[0]))
                    if not row[2]:
                        cx.execute("UPDATE contacts SET company_id=? WHERE id=?", (cid, row[0]))
                    contact_id = row[0]
                else:
                    contact_id = cx.execute(
                        "INSERT INTO contacts (email, company_id, custom_json, status, tags, created_at) VALUES (?,?,?,?,?,?)",
                        (e, cid, json.dumps(custom), status, r["source"], now)).lastrowid
                if e == r.get("email"):
                    r["contact_id"] = contact_id
    batches = []
    for src, label, rs in batch_groups(order):
        b = db.q1("SELECT id FROM batches WHERE name=?", (label,))
        if b:
            batches.append((label, len(rs), "kept"))
            continue
        bid = db.insert("batches", {"name": label, "status": "draft", "source": SOURCES[src][0],
                                    "note": f"From {SOURCES[src][2]}", "priority": SOURCES[src][1],
                                    "created_at": now})
        with db.transaction() as cx:
            for r in rs:
                if r.get("contact_id"):
                    cx.execute("INSERT OR IGNORE INTO batch_members (batch_id, contact_id, added_at) VALUES (?,?,?)",
                               (bid, r["contact_id"], now))
        batches.append((label, len(rs), "new"))
    return stats, batches


# ---------------------------------------------------------------- forms workbook

def write_forms_xlsx(order, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    F = "Arial"
    head_font = Font(name=F, bold=True, color="FFFFFF", size=10)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    input_fill = PatternFill("solid", fgColor="FFF9DB")
    body = Font(name=F, size=10)
    link = Font(name=F, size=10, color="0563C1", underline="single")
    thin = Border(bottom=Side(style="thin", color="DDDDDD"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Формы"
    cols = [("№", 6), ("Сайт поставщика", 32), ("Ссылка на форму Wholesale", 46), ("Почта", 30), ("Источник", 24),
            ("Отправили форму", 14), ("Причина почему не отправили", 28), ("Исполнитель", 14),
            ("Пароль (если форма с регистрацией)", 20), ("Комментарий / что не заполнилось", 32),
            ("Дата отправки формы", 14), ("Получен ответ (статус)", 30), ("Получен прайс-лист (ссылка)", 28),
            ("Что делала другая команда (CRM)", 40), ("Открытые оптовые цены", 22)]
    inputs = {6, 7, 8, 9, 10, 11, 12, 13}
    for i, (name, w) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    forms = [r for r in order if r["plan"] == "form"]
    forms.sort(key=lambda r: (SOURCES[r["source"]][1], r["domain"]))
    for n, r in enumerate(forms, start=1):
        c = r["crm"]
        vals = [n, r["website"] or r["domain"], r["form_url"], ", ".join(r["emails"]), SOURCES[r["source"]][0], "Нет",
                c.get("reason", ""), "", "", "", "", "", "", crm_note(r).removeprefix("Other team: "), r["price_note"]]
        for i, v in enumerate(vals, start=1):
            cell = ws.cell(row=n + 1, column=i, value=v)
            cell.font = body
            cell.border = thin
            if i in inputs:
                cell.fill = input_fill
            if i in (2, 3) and isinstance(v, str) and v.startswith("http"):
                cell.hyperlink, cell.font = v, link
        ws.cell(row=n + 1, column=11).number_format = "yyyy-mm-dd"
    last = len(forms) + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    dv_yes = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=True)
    dv_reason = DataValidation(type="list", allow_blank=True,
                               formula1='"Форма с регистрацией,Расширение не заполнило,Отправка обычной почтой,'
                                        'Не открывается сайт,Нет простой формы для регистрации,Не сотрудничают"')
    ws.add_data_validation(dv_yes)
    ws.add_data_validation(dv_reason)
    dv_yes.add(f"F2:F{last}")
    dv_reason.add(f"G2:G{last}")

    s = wb.create_sheet("Сводка")
    s.column_dimensions["A"].width = 38
    s.column_dimensions["B"].width = 14
    s["A1"], s["A1"].font = "Прогресс по формам", Font(name=F, bold=True, size=13)
    rows = [("Всего форм", f"=COUNTA(Формы!B2:B{last})"),
            ("Отправлено (Да)", f'=COUNTIF(Формы!F2:F{last},"Да")'),
            ("Осталось", "=B3-B4"),
            ("Не получилось (указана причина)", f'=COUNTIFS(Формы!F2:F{last},"Нет",Формы!G2:G{last},"<>")'),
            ("Получен прайс-лист", f'=COUNTIF(Формы!M2:M{last},"<>")')]
    for i, (k, f) in enumerate(rows, start=3):
        s.cell(row=i, column=1, value=k).font = Font(name=F, size=10)
        s.cell(row=i, column=2, value=f).font = Font(name=F, size=10, bold=True)
    s["A10"], s["A10"].font = "По источникам", Font(name=F, bold=True, size=11)
    for i, key in enumerate([k for k in SOURCES if any(r["source"] == k for r in forms)], start=11):
        label = SOURCES[key][0]
        s.cell(row=i, column=1, value=label).font = Font(name=F, size=10)
        s.cell(row=i, column=2, value=f'=COUNTIF(Формы!E2:E{last},A{i})').font = Font(name=F, size=10)
        s.cell(row=i, column=3, value=f'=COUNTIFS(Формы!E2:E{last},A{i},Формы!F2:F{last},"Да")').font = Font(name=F, size=10)
    s.cell(row=10, column=2, value="форм").font = Font(name=F, bold=True, size=10)
    s.cell(row=10, column=3, value="отправлено").font = Font(name=F, bold=True, size=10)

    g = wb.create_sheet("Инструкция")
    g.column_dimensions["A"].width = 110
    lines = [
        "Как пользоваться",
        "1. Лист «Формы» — поставщики, у которых есть wholesale-форма на сайте. Им НЕ отправляем письма из EmailManager.",
        "2. Жёлтые столбцы заполняет исполнитель: «Отправили форму» (Да/Нет), причина, исполнитель, пароль, комментарий, дата, ответ, прайс-лист.",
        "3. «Что делала другая команда (CRM)» — справочно: эти поставщики уже получали заявки от другой команды. Мы подаём свою форму заново.",
        "4. Если форму отправить невозможно (только почтой, сайт не открывается) — выберите причину; такие компании можно перенести в рассылку.",
        "5. Лист «Сводка» считает прогресс автоматически.",
        "Пример заполненной строки: Отправили форму = Да · Исполнитель = Анастасия · Дата = 2026-10-07 · Ответ = «Автоответ, ждём 3 дня».",
    ]
    for i, t in enumerate(lines, start=1):
        g.cell(row=i, column=1, value=t).font = Font(name=F, size=11, bold=(i == 1))
    wb.move_sheet("Инструкция", offset=-2)
    wb.active = 1
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(path)
    return len(forms)


def reset_import():
    """Removes draft batches and imported companies/contacts that have no real email activity."""
    with db.transaction() as cx:
        cx.execute("DELETE FROM batch_members WHERE batch_id IN (SELECT id FROM batches WHERE status='draft')")
        cx.execute("DELETE FROM batches WHERE status='draft'")
        cx.execute("DELETE FROM contacts WHERE id NOT IN (SELECT contact_id FROM enrollments) "
                   "AND id NOT IN (SELECT contact_id FROM messages WHERE contact_id IS NOT NULL) "
                   "AND id NOT IN (SELECT contact_id FROM batch_members) "
                   "AND company_id IN (SELECT id FROM companies WHERE COALESCE(plan,'')<>'')")
        cx.execute("DELETE FROM companies WHERE COALESCE(plan,'')<>'' AND id NOT IN "
                   "(SELECT company_id FROM contacts WHERE company_id IS NOT NULL) AND id NOT IN "
                   "(SELECT company_id FROM messages WHERE company_id IS NOT NULL)")


def main(argv):
    folder = Path(argv[1]) if len(argv) > 1 else Path("../ResultsCompanies")
    dry = "--dry-run" in argv
    db.init()
    companies = read_folder(folder)
    order = plan(companies)
    counts = defaultdict(lambda: defaultdict(int))
    for r in order:
        counts[SOURCES[r["source"]][0]][r["plan"]] += 1
    for src, c in counts.items():
        print(f"{src:36s} " + "  ".join(f"{k}={v}" for k, v in sorted(c.items())))
    if dry:
        for src, label, rs in batch_groups(order):
            print(f"  batch {label}: {len(rs)}")
        return
    if "--reset" in argv:
        reset_import()
    stats, batches = apply(order)
    print(dict(stats))
    for b in batches:
        print("  batch", *b)
    out = folder / "Формы для заполнения.xlsx"
    n = write_forms_xlsx(order, out)
    print(f"forms workbook: {out} ({n} rows)")


if __name__ == "__main__":
    main(sys.argv)
